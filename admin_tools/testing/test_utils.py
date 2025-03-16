# standard dependencies
import os

# 3rd-party dependencies
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils.dataframe import dataframe_to_rows

# internal dependencies
pass


def add_imgs_to_spreadsheet(df):
    csv_path = 'output/face_data.csv'
    df = pd.read_csv(csv_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "Face Data"

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=value)

    image_column = len(df.columns) + 1
    ws.cell(row=1, column=image_column, value="image")

    for idx, row in df.iterrows():
        img_path = f'output/faces/{idx}.jpg' 
        if os.path.exists(img_path):
            img = XLImage(img_path)
            img.height = 80
            img.width = 80
            ws.add_image(img, f"{chr(65 + image_column - 1)}{idx + 2}")

    output_path = 'output/face_data.xlsx'
    wb.save(output_path)

    output_path


